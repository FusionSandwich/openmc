"""Create the reviewer packet with the already bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import hashlib
import json

ROOT = Path(__file__).resolve().parent
WORKSPACE = ROOT.parents[4]
CAMPAIGN = WORKSPACE / 'stellarcsg-workload-ladder-20261008'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.04
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ('Heading 1', 'Heading 2'):
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.keep_with_next = True

def p(text, style=None):
    return doc.add_paragraph(text, style)
def h(text):
    doc.add_heading(text, 2)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def bullets(items):
    for item in items:
        p(item, 'List Bullet')
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = Inches(width)
        for r in cell.paragraphs[0].runs:
            r.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.text = text
            cell.width = Inches(width)
    for index, row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E7EDF3' if index == 0 else ('F6F7F8' if index % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '90')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    p('')

p('Fourier torus methods for stellarator plasma geometry', 'Title')
p('Independent feasibility and performance review', 'Subtitle')
p('October 9 2026\nProposal owner Joshua\nTechnical modification and context prepared by Codex')
p('Please assess two proposals separately: Joshua\'s reconstruction of a VMEC plasma boundary from simpler torus geometries, and Codex\'s use of analytic torus approximations to accelerate intersection with the original boundary. Determine whether each preserves the required geometry and particle tracking semantics, and whether it has a credible route to the performance target.')
h('The problem and the decision')
p('We need accurate neutron or photon transport through the actual nonaxisymmetric stellarator plasma boundary. The input is a VMEC wout NetCDF file; "VMC" in the request is understood as VMEC. A similarly sized built-in torus is the plasma speed reference. Complex coil packs use flat-ring references and separate geometry inputs. Plasma coefficients alone do not define coil winding packs.')
p('The target is throughput at least 80% of the reference for a frozen completed transport workload. Equivalently, custom transport time must be at most 1.25 times reference transport time. A simpler substitute boundary does not meet the geometry requirement. Construction, loading, transport and total process time must be reported separately.')
table(['Proposal', 'Role', 'Current evidence'], [
    ['Original', 'Reconstruct the stellarator using simpler torus geometries derived from its Fourier representation.', 'User concept; no conversion, completeness proof or measured speed result yet.'],
    ['Modification', 'Use torus bounds or candidates, then resolve the actual Fourier boundary.', 'Engineering hypothesis; neither speed nor global correctness established.'],
], [1.05, 3.15, 2.7])
p('The current implementations and saved results should remain the baselines. This packet requests an independent assessment, not endorsement of either proposal. Review mathematical feasibility before a broad implementation or transport campaign.')

page('The original proposal')
p('Joshua described the method as follows. This quotation preserves the original proposal before technical interpretation.')
p('"We take the Fourier transform of the plasma geometry that we get from the VMEC file, and then what we do is we transfer that into a phase space that allows us to break it down into individual toruses. So we treat the toruses, we treat a cosine function that would be regularly used for a Fourier transform. And so what we do is we take the complex geometry, we can break it down into simpler torus geometry, and then we can combine those together to fully create the complex stellarator geometry while using the simpler toroidal CSG geometry that we have analytical solutions for."')
h('Proposed sequence')
bullets(['Read the plasma geometry from VMEC and obtain its Fourier representation.', 'Transform that representation into a space in which elementary toroidal geometries act as basis components.', 'Reconstruct the complex plasma using those components.', 'Use analytic torus intersections instead of the expensive general boundary solver.'])
h('Definitions the reviewer must resolve')
p('The proposal does not yet specify the transform, its inverse, the torus parameters, or the composition operation. Assess separately: a Boolean union or intersection of torus solids; a weighted sum of torus implicit functions; a sum of parametric coordinate displacements; or a nonlinear coordinate map to a reference torus. These interpretations have different boundary equations and tracking costs. Do not silently replace the proposal with one interpretation and claim to have tested them all.')
p('VMEC already supplies Fourier coefficients. For a stellarator-symmetric boundary, R(theta, phi) = sum R_mn cos(m theta - n Nfp phi), and Z(theta, phi) = sum Z_mn sin(m theta - n Nfp phi). Nonsymmetric files also need the complementary sine and cosine coefficients. In a wout file xn already incorporates Nfp; multiplying it again is an error. Coordinates are in metres and the current transport reader converts them to centimetres. [1, 3]')
h('A concrete feasibility test')
p('Consider R = R0 + [a0 + epsilon cos(N phi)] cos(theta), Z = [a0 + epsilon cos(N phi)] sin(theta). Its tube radius varies around the ring. A Fourier harmonic changes the same surface; it is not automatically an independent solid torus. Boolean combinations of concentric tori sharing the same axis retain rotational symmetry, so that restricted construction cannot reproduce this deformation. Differently placed or trimmed tori could approximate it, but need a fitting and error analysis.')
p('A nonlinear map may simplify the boundary while turning a straight physical ray into a curved path in transformed coordinates. Review whether solving that path recovers the original cost. These are concerns to investigate; no general impossibility theorem for all torus-based constructions is claimed.')

page('The suggested modification')
p('Preserve the actual VMEC boundary as the geometric target. Introduce one analytic reference torus or a small atlas of local torus patches to reduce the work needed to locate the first physical intersection. Analytic roots are candidates or bounds, rather than automatically accepted plasma crossings.')
h('Geometry representation')
p('Let X(theta, phi) denote the boundary reconstructed from all admitted VMEC coefficients. Fit T(theta, phi), or local T_j, and retain a residual D = X - T over the same parameter domains. For each domain obtain conservative bounds on position error, derivative error where required, and the spatial region containing that boundary patch. A small sampled residual alone is insufficient for whole-domain exclusion.')
p('For omitted Fourier terms, sum absolute coefficient magnitudes to bound each coordinate residual uniformly. Combining radial and vertical coordinate bounds gives a Euclidean pointwise bound for matching parameters. That bound alone does not prove signed-distance, normal, volume or ray-intersection error. Normal bounds also require derivative bounds and a lower bound on a nondegenerate surface Jacobian. Hausdorff closeness alone does not license inside/outside tests.')
h('Proposed query path')
bullets(['Accept a physical origin, direction, coincidence state and optional collision-distance ceiling.', 'Use existing conservative component or patch bounds and clearance certificates to exclude irrelevant domains.', 'Evaluate the reference torus or relevant local torus candidates and combine them with validated residual enclosures.', 'Resolve candidate intervals against the actual boundary using bounded refinement or an independently checked solver.', 'Exclude every earlier possible interval before returning the first crossing. Handle grazing contacts, seams and coincident starts explicitly.', 'Return HIT, NO HIT or UNRESOLVED. Failure to find a numerical seed is not proof of a miss. Preserve normals, Boolean sense and successor ownership.'])
h('Important alternatives')
p('A global torus fit may be too loose for a strongly shaped stellarator. A local atlas may provide tighter bounds but costs more traversal, fitting and storage. A direct Fourier evaluator may avoid spline fitting error while costing more trigonometric evaluations. Compare both where practical, and separate representation changes from accelerator effects.')
p('The current spline payload is an approximation derived from VMEC, whereas a new Fourier backend may use the retained modes directly. Bitwise agreement is expected only for semantics-preserving acceleration of the same representation. A changed representation needs an independent geometric and scoring comparison, not a claim that tally differences alone prove failure or success.')

page('Current methods and measured context')
p('The following results are saved campaign observations as inspected on October 9 2026. The campaign is owned separately by Claude. Its checkout and processes must not be modified. Saved passing gates do not by themselves close independent physical geometry qualification. [4]')
table(['Method or experiment', 'Result', 'Interpretation'], [
    ['P00 faceted coil proxy with combined optimizations', 'gate 03 ratio 1.1046; about 90.5% reference throughput. All three seeds PASS ROBUST.', 'Represented mesh versus fixed 30 by 30 cm rings. Continuous CAD, physical ownership and independent qualification remain open.'],
    ['Plasma collision-distance ceiling', '8.07 times faster than control. Point-source custom/torus ratio 195.9 to 23.8.', 'Source starts inside custom iron and outside torus; 12.6 times collision workload mismatch. Not a fair geometry-only comparison.'],
    ['Plasma comparable-workload source diagnostic', 'v2 to v7 ratio 2.100; v2 to v8 ratio 2.120. About 47 to 48% reference throughput.', 'Uniform source inside each plasma; 51.6 versus 60.6 collisions per source. Diagnostic, different boundaries and sources; not a final matched acceptance gate.'],
    ['Voxel clearance ahead of BVH certificate', '1.139 times median speedup; uncertainty band 1.092 to 1.216. Initialization adds 0.39 s.', 'Additional effect undecided against the registered conservative 1.10 criterion.'],
    ['Anchored native axis events', '39 models, 208 rows; five deliberate corruptions rejected.', 'Accepted restricted deterministic void unit-axis plane slice. No curved-plasma or speed qualification.'],
], [2.0, 2.15, 2.75])
p('The comparable-workload ratio 2.12 would require about 41% less custom transport time to reach 1.25, holding the reference fixed. The latest matched plasma profile attributes 44.8% of time to the full solver and 13.2% to exclusion certificates. About 92% of distance calls are answered by certificates; the expensive residual calls still matter. These figures motivate targeting surviving solver work, but do not predict a torus-based speedup. [4]')
p('Historical fast swept-coil solvers retain wrong-first-root counterexamples. The global offset and interval-cache method is a restricted correctness reference with costly transport. Exact torus dispatch is valuable for genuinely analytic inputs, but cannot substitute for the nonaxisymmetric target. No qualified common-workload comparison across every historical method exists. [5]')

page('Independent review and acceptance questions')
h('Mathematical feasibility')
bullets(['Original proposal: define the basis, transform, reconstruction operation and inverse. Is the result exact for the retained VMEC modes, or an approximation? What class of boundaries is admitted?', 'If using a transformed coordinate space, derive physical ray tracking and Euclidean flight length. Explain normals and scattering directions under the map.', 'If using Boolean torus primitives, quantify required primitive count versus spatial tolerance, and handle trims, internal seams and duplicate contacts.', 'Modification: show conservative residual and exclusion bounds over entire domains, complete earliest-root coverage and explicit unresolved behavior.', 'Assess validity for nonaxisymmetry, field periodicity, asymmetric Fourier terms, noncircular sections and coordinate singularities.'])
h('Geometry and tracking tests')
p('Begin with exact circular and elliptical tori, one known nonaxisymmetric harmonic, a multimode surface, and the provenance-bound WISTELL VMEC boundary. Check units, phase conventions, field-period scaling and reproducibility. Freeze an independent bank containing ordinary misses, competing first roots, near tangencies, even contacts, thin chords, seam wraps, coincidence and transformed placements. A converged root is insufficient without earlier-root exclusion.')
p('Compare position, normals, signed membership, first entry and exit distances, and successor ownership. Use independent high-precision or interval reference calculations for difficult cases; the existing checked spline solver is a comparator whose general completeness also has open obligations. Specify tolerances from the project accuracy budget before fitting or timing. This packet does not invent a physical tolerance.')
h('Performance and native scoring')
p('Measure fit and construction time, memory, loading, certificate hit rate, surviving candidates, iterations, fallback rate and uninstrumented query time. Then compare frozen native transport with common physics and tally settings, sources with comparable occupancy and collision work, several seeds and useful run durations. Preserve the existing minimum 20k histories and registered repeatability rules. Compare local tally bins, not only global sums; inspect currents and ownership as well as track length.')
p('Use an ablation matrix: incumbent; torus assistance alone; incumbent accelerators alone; and their combination. Include the same torus reference runtime. A faster query microbenchmark is not proof of 80% transport throughput. Report speed versus the incumbent separately from speed relative to the native torus.')
h('Requested reviewer output')
p('For each proposal return FEASIBLE AS STATED, FEASIBLE WITH SPECIFIED CHANGES, UNSUPPORTED or COUNTEREXAMPLE, with assumptions and decisive reasoning. Separately report correctness status, geometry fidelity, expected cost scaling, implementation effort and the smallest experiment that could change the conclusion. Do not assign a numerical speedup without measurements or a justified cost model.')

page('Implementation sequence and compatible methods')
p('Start in a new local Codex chat with fresh history and a unique prototype directory. This chat continues the existing combination work in a different directory. Use only existing local capabilities initially; no new dependencies, remote jobs or broad overlapping campaigns are required for the first mathematical prototype.')
table(['Stage', 'Concrete output', 'Decision before proceeding'], [
    ['1  Define and test the original interpretation', 'Explicit torus composition API and reconstruction test on exact torus and one harmonic.', 'Record what succeeds, fails or remains undefined. Keep the user proposal separate from the modification.'],
    ['2  Fit and bound torus assistance', 'VMEC loader, reference torus fit, residual bounds and candidate interval prototype.', 'Bounds remain conservative on synthetic adversarial examples; unsupported domains explicit.'],
    ['3  Evaluate combinations', 'Independent ablation harness for bounds, ceiling, clearance and solver candidates.', 'Accuracy and first-root obligations satisfied before attributing timing benefit.'],
    ['4  Native integration', 'Opt-in backend or accelerator in an isolated OpenMC worktree.', 'Independent review, finite root bank, native owner and per-bin scoring acceptance.'],
    ['5  Performance qualification', 'Repeated same-build incumbent/candidate/reference transport receipts.', 'Total transport ratio at most 1.25 plus geometry and workload qualification.'],
], [1.45, 2.95, 2.5])
h('Methods that may compose')
p('Conservative BVHs and local-span hierarchies, collision-distance ceilings, voxel clearance and faster exact arithmetic operate at different stages and can potentially be retained. Torus seeds add another candidate generator. A validated cache must key all state affecting the query, including geometry identity, origin, direction, coincidence and ceiling where relevant. Interval or Bernstein methods may resolve difficult candidate intervals; their admitted geometry and completeness contracts must carry over.')
p('No accelerator may silently turn an unresolved query into NO HIT. The axis-plane event scaffold is restricted to its tested domain and cannot be promoted to curved trajectories by analogy. Speed factors are not automatically multiplicative because methods can eliminate the same work.')
h('Isolation and publication')
p('Do not edit or switch the Claude checkout stellarcsg-astra-kernel-20260926, alter its campaign inputs or libraries, control its processes or take over its compute ownership. Read its receipts only. Use disjoint source and output directories. Record input, source and runtime hashes. Before compilation or acquisition, collect the current operator inventory and obey resource leases; avoid simultaneous heavy campaigns.')
p('Publish new reviewable source and compact evidence on separate JS branches with draft PRs. Preserve the current axis integration and measured baselines. Do not push another agent\'s uncommitted work, change its PR or merge without the required review. The user explicitly authorized branches and PRs on October 9 2026.')

page('Sources and evidence bindings')
p('The accompanying source manifest records SHA256 hashes of the local files used to prepare this packet. Paths below are relative to the StellarCGS workspace unless otherwise stated. A review using only this document is conceptual; execution and raw-array validation require the bound artifacts.')
h('External primary references')
p('[1] Simsopt SurfaceRZFourier documentation. Fourier boundary equations, stellarator symmetry, VMEC imports and surface parameterization.\nhttps://simsopt.readthedocs.io/stable/simsopt.geo.html#simsopt.geo.surfacerzfourier.SurfaceRZFourier')
p('[2] OpenMC geometry methods. CSG composition, physical ray distance and torus quartic solver.\nhttps://docs.openmc.org/en/latest/methods/geometry.html#torus-parallel-to-an-axis')
h('Local representation and acceptance context')
p('[3] stellarcsg-axis-integration-20261008/dev/stellarcsg/python/stellarcsg/vmec.py\nExisting VMEC coefficient reader; position and derivative reconstruction.\nSame checkout: dev/stellarcsg/include/stellarcsg/periodic_radial_surface.hpp\nHIT, NO HIT and UNRESOLVED interface.')
p('[4] stellarcsg-workload-ladder-20261008/RESULTS.md and PLASMA_PLAN.md\nRaw summaries: runs/gate-03/GATE.json; plasma/runs/plasma-ab-01/SUMMARY.json; plasma/runs/plasma-matched-01/SUMMARY.json; plasma/runs/plasma-matched-02/SUMMARY.json.\nThese bind the numerical context reported on page 4. Newer campaign updates must be treated as new evidence rather than silently replacing these snapshots.')
p('[5] stellarcsg-axis-integration-20261008/dev/stellarcsg/METHOD_MAP.md and dev/stellarcsg/reports/BEST_METHOD_CONTINUATION_20260930.md\nHistorical method roles, failures and scope. Numerical summaries predate the October 9 campaign and are historical context.')
p('[6] stellarcsg-axis-integration-20261008/INTEGRATION_RESULT_20261008.json and integration-20261008/independent-final-review.json\nAccepted axis slice and noninterference evidence; no plasma speed claim.')
p('[7] C:/Users/joshu/Downloads/StellarCSG_review_findings_and_test_contract.md and USABILITY_METHOD_REVIEW_20260926.md\nIndependent static review and earlier method review. Their recommendations are source material, distinct from the user\'s request and this proposed experiment.')
h('Scope of this packet')
p('The original proposal is quoted from the user\'s October 9 message. The modified method is Codex\'s suggestion. No implementation or timing of either new method existed at packet creation. Current measured speed evidence belongs to separate incumbent experiments. An independent reviewer should challenge both proposals and may recommend a third approach if supported by a concrete argument.')

footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('Page ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Fourier torus methods for stellarator plasma geometry'
doc.core_properties.author = 'Joshua and Codex'
for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(ROOT / 'Fourier_torus_methods_reviewer_packet.docx')
inputs = [CAMPAIGN / 'RESULTS.md', CAMPAIGN / 'PLASMA_PLAN.md',
          CAMPAIGN / 'runs/gate-03/GATE.json',
          CAMPAIGN / 'plasma/runs/plasma-ab-01/SUMMARY.json',
          CAMPAIGN / 'plasma/runs/plasma-matched-01/SUMMARY.json',
          CAMPAIGN / 'plasma/runs/plasma-matched-02/SUMMARY.json',
          ROOT.parents[3] / 'dev/stellarcsg/python/stellarcsg/vmec.py',
          ROOT.parents[3] / 'INTEGRATION_RESULT_20261008.json',
          ROOT.parents[3] / 'dev/stellarcsg/include/stellarcsg/periodic_radial_surface.hpp',
          ROOT.parents[3] / 'dev/stellarcsg/METHOD_MAP.md',
          ROOT.parents[3] / 'dev/stellarcsg/reports/BEST_METHOD_CONTINUATION_20260930.md',
          ROOT.parents[3] / 'integration-20261008/independent-final-review.json',
          Path('C:/Users/joshu/Downloads/StellarCSG_review_findings_and_test_contract.md'),
          Path('C:/Users/joshu/Downloads/USABILITY_METHOD_REVIEW_20260926.md')]
manifest = {'date': '2026-10-09', 'files': [], 'new_method_status': 'CONCEPT_NOT_YET_MEASURED'}
snapshot_dir = ROOT / 'evidence_snapshots'
snapshot_dir.mkdir(exist_ok=True)
for index, path in enumerate(inputs):
    data = path.read_bytes()
    snapshot = snapshot_dir / (str(index + 1) + '_' + path.name)
    snapshot.write_bytes(data)
    manifest['files'].append({'source': str(path), 'snapshot': str(snapshot.relative_to(ROOT)),
                              'sha256': hashlib.sha256(data).hexdigest()})
(ROOT / 'source_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
print(ROOT / 'Fourier_torus_methods_reviewer_packet.docx')
